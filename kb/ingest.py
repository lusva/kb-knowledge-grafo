"""Estrazione testo dai documenti e chunking."""
import html as _html
import logging
import re
import tempfile
import time
from pathlib import Path

from django.conf import settings

from . import config, jobs

ingest_log = logging.getLogger("kb.ingest")


def html_to_text(html: str) -> str:
    """Riduce l'HTML a testo leggibile: via script/style, tag ed entità.

    È la stessa pulizia per i corpi HTML delle email (Gmail, Microsoft 365,
    .eml, .pst) e per le pagine .html importate come documenti."""
    text = re.sub(r"(?is)<(script|style|head)[^>]*>.*?</\1>", " ", html)
    text = re.sub(r"(?i)<br\s*/?>|</p>|</div>|</tr>|</li>|</h[1-6]>", "\n", text)
    text = re.sub(r"<[^>]+>", " ", text)
    text = _html.unescape(text)
    text = re.sub(r"[ \t\xa0]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)      # niente spazi a filo di riga
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def extract_text(path: Path, ocr_max_pages: int = None,
                 ocr_max_seconds: float = None, ocr_max_chars: int = None) -> str:
    """Estrae il testo grezzo da un file (non immagini).

    ``ocr_max_pages`` limita le pagine sottoposte a OCR per i PDF scansionati
    (default: KB_OCR_MAX_PAGES). Serve agli allegati della chat, dove la
    risposta deve arrivare in tempi da richiesta web. ``ocr_max_chars`` ferma
    l'OCR quando il testo letto basta (la lettura delle schede, che del
    documento usa solo quel che sta nel contesto del modello)."""
    ext = path.suffix.lower()
    if ext in settings.KB_TEXT_EXT:
        text = path.read_text(encoding="utf-8", errors="replace")
        return text
    if ext in settings.KB_PDF_EXT:
        return pdf_text(path, ocr_max_pages, ocr_max_seconds, ocr_max_chars)
    if ext in settings.KB_DOCX_EXT:
        import docx
        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    if ext in settings.KB_PPTX_EXT:
        from pptx import Presentation
        prs = Presentation(str(path))
        parts = []
        for i, slide in enumerate(prs.slides, 1):
            parts.append(f"[Slide {i}]")
            for shape in slide.shapes:
                if shape.has_text_frame:
                    parts.append(shape.text_frame.text)
        return "\n".join(parts)
    if ext in settings.KB_XLSX_EXT:
        return xlsx_text(path)
    if ext in settings.KB_HTML_EXT:
        return html_to_text(path.read_text(encoding="utf-8", errors="replace"))
    if ext in settings.KB_EML_EXT:
        from . import eml
        return eml.extract_text(path)
    if ext in settings.KB_CAD_EXT:
        from . import cad
        return cad.dxf_text(path)
    raise ValueError(f"Estensione non supportata: {ext}")


# --------------------------------------------------------------------------- #
# PDF (PyMuPDF): testo con l'ordine di lettura della pagina, e per le
# scansioni la rasterizzazione delle pagine senza programmi esterni.
# --------------------------------------------------------------------------- #

def _pymupdf():
    """Il modulo PyMuPDF, con i messaggi diagnostici spenti: su un PDF un po'
    rotto MuPDF stampa una riga di avviso per ogni oggetto storto, e il log
    dell'indicizzazione si riempie di "invalid ICC colorspace" innocui."""
    import pymupdf
    pymupdf.TOOLS.mupdf_display_errors(False)
    return pymupdf


def pdf_text(path: Path, ocr_max_pages: int = None,
             ocr_max_seconds: float = None, ocr_max_chars: int = None) -> str:
    pymupdf = _pymupdf()

    with pymupdf.open(str(path)) as doc:
        if doc.needs_pass:
            doc.authenticate("")      # cifrato con password vuota: si apre lo stesso
        text = "\n\n".join(page.get_text("text") for page in doc)
    # PDF fatto di pagine-immagine (scansione): non c'è testo da leggere, va
    # ricavato con l'OCR. Vale anche per i PDF il cui "testo" è una sfilza di
    # codici di glifo senza mappa Unicode (\x00\x01\x02…): sono caratteri, ma
    # non si leggono, e contarli come testo lascerebbe il documento vuoto.
    if testo_illeggibile(text):
        return ocr_pdf(path, ocr_max_pages, ocr_max_seconds, ocr_max_chars) or text
    return text


def caratteri_leggibili(testo: str) -> int:
    """Quante lettere e cifre contiene il testo: la misura per decidere se un
    PDF ha davvero un livello di testo o se va passato all'OCR."""
    return sum(1 for c in testo if c.isalnum())


QUOTA_RUMORE = 0.3   # oltre questa frazione di caratteri di controllo, il testo è rumore


def _rumore(c: str) -> bool:
    """Carattere che in un testo vero non compare: di controllo (non spazio),
    area a uso privato dei font, segnaposto di decodifica fallita."""
    o = ord(c)
    return (o < 32 and not c.isspace()) or 0xE000 <= o <= 0xF8FF or o == 0xFFFD


def testo_illeggibile(testo: str) -> bool:
    """True se il livello di testo del PDF non serve: troppo corto, oppure
    pieno di codici di glifo (\\x00\\x01\\x02…) al posto delle lettere. I
    moduli con righe di puntini e trattini restano testo: la punteggiatura
    non è rumore."""
    leggibili = caratteri_leggibili(testo)
    if leggibili < int(config.get("ocr_min_chars")):
        return True
    pieni = sum(1 for c in testo if not c.isspace())
    rumore = sum(1 for c in testo if _rumore(c))
    return pieni > 0 and rumore / pieni > QUOTA_RUMORE


def ocr_pdf(path: Path, max_pages: int = None, max_seconds: float = None,
            max_chars: int = None) -> str:
    """Legge un PDF scansionato: rasterizza le pagine e le passa al modello OCR.

    Le pagine vengono rese in scala di grigi e a ``ocr_width`` pixel di
    larghezza: oltre non si guadagna accuratezza e il modello diventa molto più
    lento. Ritorna stringa vuota se l'OCR è disattivato o non produce nulla."""
    pymupdf = _pymupdf()

    from . import ollama

    if not config.get("ocr_model"):
        return ""
    limite = max_pages or settings.KB_OCR_MAX_PAGES
    # Una pagina scansionata può tenere occupato il modello per minuti, e il
    # numero di pagine da solo non dice quanto durerà il lavoro. Chi chiama
    # dentro una richiesta HTTP passa un tetto di tempo: quel che si è letto
    # fin lì si tiene, il resto si lascia.
    scadenza = (time.monotonic() + max_seconds) if max_seconds else None
    pagine, letti = [], 0
    with tempfile.TemporaryDirectory() as tmp, pymupdf.open(str(path)) as doc:
        if doc.needs_pass:
            doc.authenticate("")
        for numero, page in enumerate(doc, 1):
            jobs.controlla()      # fermata dall'utente: non aspettare la fine del PDF
            if numero > limite:
                break
            if scadenza and time.monotonic() > scadenza:
                break
            if max_chars and letti >= max_chars:
                break
            larghezza = page.rect.width or 595
            zoom = min(config.get("ocr_width") / larghezza, settings.KB_OCR_DPI / 72)
            pix = page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom),
                                  colorspace=pymupdf.csGRAY)
            img = Path(tmp) / f"pag-{numero}.png"
            pix.save(str(img))
            try:
                testo = ollama.ocr_image(img)
            except ollama.OllamaIrraggiungibile:
                raise
            except ollama.OllamaError as e:
                # una pagina che il modello non riesce a leggere non deve
                # costare le altre ventotto: si segna e si va avanti
                ingest_log.warning("OCR di %s, pagina %d: %s", path.name, numero, e)
                testo = f"[pagina {numero} non letta]"
            if testo:
                pagine.append(testo)
                letti += len(testo)
    return "\n\n".join(pagine)


# Il lato corto di un A4 è 595 punti, quello di un A3 842: da A3 in su una
# pagina di PDF è quasi sempre una tavola (disegno tecnico, schema, layout).
# Lì il livello di testo non aiuta: sono quote sparse senza ordine, e nelle
# tavole scansionate spesso un OCR fatto male da chi le ha esportate, che passa
# per testo buono (433.03.005.0.pdf: 1600 caratteri di "C)", "+I", "-.;;t").
LATO_MIN_DISEGNO = 780


def pagina_disegno(page) -> bool:
    return min(page.rect.width, page.rect.height) >= LATO_MIN_DISEGNO


def immagine_tavola(pymupdf, page, dest: Path) -> Path:
    """Salva la pagina come PNG con il lato lungo a ``disegni_width`` pixel:
    la risoluzione a cui il modello dei disegni legge la tavola."""
    zoom = config.get("disegni_width") / max(page.rect.width, page.rect.height)
    page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom)).save(str(dest))
    return dest


def descrivi_disegni_pdf(path: Path, max_pages: int, max_seconds: float = None) -> str:
    """Se il PDF contiene tavole, le fa descrivere al modello di visione;
    le pagine normali dello stesso PDF restano al loro testo.

    Ritorna "" se non ci sono tavole o se nessuna è stata descritta: chi
    chiama ripiega sull'estrazione normale."""
    pymupdf = _pymupdf()

    from . import ollama

    scadenza = (time.monotonic() + max_seconds) if max_seconds else None
    parti, descritte = [], 0
    with tempfile.TemporaryDirectory() as tmp, pymupdf.open(str(path)) as doc:
        if doc.needs_pass:
            doc.authenticate("")
        if not any(pagina_disegno(p) for p in doc):
            return ""
        for numero, page in enumerate(doc, 1):
            jobs.controlla()
            if numero > max_pages or (scadenza and time.monotonic() > scadenza):
                parti.append(f"[… altre {len(doc) - numero + 1} pagine non lette]")
                break
            if not pagina_disegno(page):
                parti.append(f"[Pagina {numero}]\n{page.get_text('text').strip()}")
                continue
            img = immagine_tavola(pymupdf, page, Path(tmp) / f"tavola-{numero}.png")
            desc = ollama.describe_image(img, ollama.DISEGNO_PROMPT, model=ollama.modello_disegni(),
                                         num_ctx=int(config.get("disegni_num_ctx")))
            if desc:
                descritte += 1
                parti.append(f"[Pagina {numero} · disegno tecnico, descritto dal modello di visione]\n{desc}")
    return "\n\n".join(parti) if descritte else ""


# --------------------------------------------------------------------------- #
# Fogli di calcolo
# --------------------------------------------------------------------------- #

MAX_RIGHE_XLSX = 5000   # oltre, il foglio è un database, non un documento


def xlsx_text(path: Path) -> str:
    """Ogni riga non vuota del foglio diventa una riga di testo, celle
    separate da " | "; ogni foglio comincia con il proprio nome."""
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    parti, righe = [], 0
    try:
        for ws in wb.worksheets:
            parti.append(f"[Foglio: {ws.title}]")
            for row in ws.iter_rows(values_only=True):
                celle = [str(c).strip() for c in row if c is not None and str(c).strip()]
                if not celle:
                    continue
                parti.append(" | ".join(celle))
                righe += 1
                if righe >= MAX_RIGHE_XLSX:
                    parti.append("[… foglio troncato]")
                    return "\n".join(parti)
    finally:
        wb.close()
    return "\n".join(parti)


def _cella(v):
    """Il valore di una cella in JSON: numeri come numeri, date come testo ISO."""
    if v is None or isinstance(v, (bool, int, str)):
        return v
    if isinstance(v, float):
        return v if v == v and v not in (float("inf"), float("-inf")) else None
    if hasattr(v, "isoformat"):
        return v.isoformat()
    try:
        return float(v)                       # Decimal
    except (TypeError, ValueError):
        return str(v)


def celle_foglio(path: Path) -> dict | None:
    """Le celle di un foglio di calcolo, per i conti con un programma
    (fogli.py): {"fogli": [{"nome", "righe": [[cella, …], …]}]}. A
    differenza di xlsx_text le posizioni restano: righe[0] è la riga 1, la
    cella i è la colonna i (A = 0), le vuote sono None. Un CSV è un foglio
    solo, con le celle come testo. None se non è un foglio di calcolo."""
    ext = path.suffix.lower()
    if ext == ".csv":
        import csv
        testo = path.read_text(encoding="utf-8", errors="replace")
        try:
            dialetto = csv.Sniffer().sniff(testo[:5000], delimiters=",;\t|")
        except csv.Error:
            dialetto = csv.excel
        import itertools
        righe = list(itertools.islice(csv.reader(testo.splitlines(), dialetto), MAX_RIGHE_XLSX))
        return {"fogli": [{"nome": path.stem, "righe": righe}]}
    if ext not in settings.KB_XLSX_EXT:
        return None
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    fogli, quante = [], 0
    try:
        for ws in wb.worksheets:
            righe = []
            for row in ws.iter_rows(values_only=True):
                riga = [_cella(c) for c in row]
                while riga and riga[-1] in (None, ""):
                    riga.pop()
                righe.append(riga)
                quante += 1
                if quante >= MAX_RIGHE_XLSX:
                    break
            while righe and not righe[-1]:
                righe.pop()
            fogli.append({"nome": ws.title, "righe": righe})
            if quante >= MAX_RIGHE_XLSX:
                break
    finally:
        wb.close()
    return {"fogli": fogli}


def extract_wikilinks(text: str) -> list[str]:
    """Trova i link stile Obsidian [[Nome nota]]."""
    return [m.strip() for m in re.findall(r"\[\[([^\]|#]+)(?:[|#][^\]]*)?\]\]", text)]


# Un tag fatto di sole cifre esadecimali è quasi sempre un colore CSS
# (#ffffff, #0b0e11) pescato dallo stile di una email HTML: non è un
# argomento, e nel grafo diventa un nodo che non vuol dire niente.
_COLORE = re.compile(r"^[0-9a-f]{3,8}$", re.I)


def extract_tags(text: str) -> list[str]:
    """Trova i tag #stile-obsidian nel testo."""
    tags = re.findall(r"(?:^|\s)#([\w/\-]+)", text)
    # esclude ancore markdown tipo "# Titolo" (già gestite dal \s ma per sicurezza)
    return sorted({t.lower() for t in tags
                   if not t.isdigit() and not _COLORE.match(t) and len(t) > 1})


def chunk_text(text: str, size: int = None, overlap: int = None) -> list[str]:
    """Divide il testo in blocchi di ~size caratteri, spezzando sui paragrafi
    quando possibile, con sovrapposizione tra blocchi consecutivi."""
    size = size or config.get("chunk_size")
    overlap = overlap or config.get("chunk_overlap")
    # Rete di sicurezza: sotto, il taglio duro avanza di `size - overlap`
    # caratteri per giro. Se la sovrapposizione arriva alla dimensione del
    # chunk quel passo è zero, il paragrafo non si accorcia mai e il ciclo non
    # finisce più. La configurazione lo impedisce (config.LIMITI), ma qui
    # arrivano anche i valori passati dal codice.
    size = max(int(size), 1)
    overlap = min(max(int(overlap), 0), size - 1)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if not text:
        return []
    if len(text) <= size:
        return [text]

    paragraphs = text.split("\n\n")
    chunks, current = [], ""
    for para in paragraphs:
        # paragrafo singolo più lungo di size: taglio duro
        while len(para) > size:
            if current:
                chunks.append(current.strip())
                current = ""
            chunks.append(para[:size])
            para = para[size - overlap:]
        if len(current) + len(para) + 2 > size and current:
            chunks.append(current.strip())
            current = current[-overlap:] if overlap else ""
        current = (current + "\n\n" + para).strip()
    if current.strip():
        chunks.append(current.strip())
    return [c for c in chunks if c.strip()]


# --------------------------------------------------------------------------- #
# Chunk dei fogli di calcolo
#
# Il testo di un foglio (xlsx_text) è una riga per riga del foglio, senza
# paragrafi: chunk_text lo tagliava ogni chunk_size caratteri anche a metà di
# una riga, e un chunk a cavallo di due fogli portava il totale di agosto
# subito prima di «[Foglio: Settembre]». Il modello lo attribuiva a settembre.
# Qui ogni foglio ha i suoi chunk, fatti di righe intere, e ognuno comincia
# con il nome del foglio.
# --------------------------------------------------------------------------- #

_FOGLIO = re.compile(r"^\[Foglio: (.*)\]$", re.M)


def dividi_fogli(text: str) -> list[tuple[str, str]]:
    """Il testo di un foglio di calcolo (xlsx_text) diviso per foglio:
    [(nome, righe)], senza i fogli vuoti. [] se il testo non ha fogli."""
    segni = list(_FOGLIO.finditer(text or ""))
    fogli = []
    for i, m in enumerate(segni):
        fine = segni[i + 1].start() if i + 1 < len(segni) else len(text)
        righe = text[m.end():fine].strip("\n")
        if righe.strip():
            fogli.append((m.group(1), righe))
    return fogli


def chunk_foglio(nome: str, righe: str, size: int = None, overlap: int = None) -> list[str]:
    """I chunk di un foglio: righe intere fino a ``size`` caratteri, ognuno
    con «[Foglio: nome]» in testa; fra un chunk e il successivo si ripetono
    le ultime righe, fino a ``overlap`` caratteri. Una riga più lunga di un
    chunk si taglia."""
    testa = f"[Foglio: {nome}]"
    size = max(int(size or config.get("chunk_size")) - len(testa) - 1, 40)
    overlap = min(max(int(overlap if overlap is not None else config.get("chunk_overlap")), 0),
                  size // 2)
    chunks, correnti, lunghezza = [], [], 0
    for riga in righe.splitlines():
        if not riga.strip():
            continue
        pezzi = [riga[i:i + size] for i in range(0, len(riga), size)]
        for pezzo in pezzi:
            if correnti and lunghezza + len(pezzo) + 1 > size:
                chunks.append(correnti)
                # sovrapposizione: le ultime righe intere che stanno in overlap
                ripresi, peso = [], 0
                for r in reversed(correnti):
                    if peso + len(r) + 1 > overlap:
                        break
                    ripresi.insert(0, r)
                    peso += len(r) + 1
                correnti, lunghezza = ripresi, peso
            correnti.append(pezzo)
            lunghezza += len(pezzo) + 1
    if correnti:
        chunks.append(correnti)
    return [testa + "\n" + "\n".join(c) for c in chunks]


def chunk_documento(text: str) -> list[str]:
    """I chunk di un documento: per foglio se è un foglio di calcolo
    (dividi_fogli), altrimenti chunk_text."""
    fogli = dividi_fogli(text)
    if not fogli:
        return chunk_text(text)
    return [c for nome, righe in fogli for c in chunk_foglio(nome, righe)]


# --------------------------------------------------------------------------- #
# Intestazione di contesto
#
# Un chunk preso a metà di un documento non sa da dove viene: "come da
# accordi, confermiamo le quantità" senza sapere che è l'email di Rossi del 3
# maggio sull'ordine 4471. Aggiungere una riga con titolo, autore e data a
# ogni pezzo prima di calcolarne il vettore rende ogni chunk cercabile da
# solo — e passa la stessa informazione al modello che lo legge.
# --------------------------------------------------------------------------- #

MAX_INTESTAZIONE = 160
# La cartella va in fondo e non si taglia mai (al più se ne tiene la parte
# finale, la più precisa): si accorcia il resto. Così una riga già scritta si
# riconosce dalla fine (ha_cartella) e il titolo lungo di un file non la perde.
MAX_CARTELLA = 90
ETICHETTE = {"email": "Email", "immagine": "Immagine", "archivio": "Archivio di posta",
             "database": "Database", "evento": "Evento"}
_INTESTAZIONE = re.compile(r"^\[(Documento|%s)\b.*\]$" % "|".join(ETICHETTE.values()))


def _coda(cartella: str) -> str:
    return f" · in {cartella[-MAX_CARTELLA:]}" if cartella else ""


def _chiudi(riga: str, cartella: str = "") -> str:
    coda = _coda(cartella)
    return f"[{riga[:MAX_INTESTAZIONE - len(coda)].rstrip()}{coda}]"


def intestazione_chunk(titolo: str, autore: str = "", data=None, tipo: str = "",
                       cartella: str = "") -> str:
    """La riga di contesto da mettere in testa a ogni chunk del documento.
    ``cartella``: dove l'utente ha messo il file, «Documenti/Fatture/fatture
    attive» (vedi reindex.cartella_del_documento)."""
    etichetta = ETICHETTE.get(tipo, "Documento")
    parti = [f"{etichetta}: {titolo.strip()}"] if titolo and titolo.strip() else [etichetta]
    if autore and autore.strip():
        parti.append(f"da {autore.strip()}")
    if data:
        parti.append(data.strftime("%d/%m/%Y"))
    return _chiudi(" · ".join(parti), cartella)


def e_intestazione(riga: str) -> bool:
    """Se la riga è un'intestazione di chunk (e non, per esempio, «[Foglio: …]»)."""
    return bool(_INTESTAZIONE.match(riga or ""))


def ha_cartella(intestazione: str, cartella: str) -> bool:
    return intestazione.endswith(_coda(cartella) + "]")


def con_cartella(intestazione: str, cartella: str) -> str:
    """Un'intestazione già scritta, con la cartella in fondo: per i documenti
    indicizzati prima che la cartella ci fosse, senza rileggerli."""
    return _chiudi(intestazione.strip()[1:-1], cartella)


def con_intestazione(chunks: list[str], intestazione: str) -> list[str]:
    """I chunk con la riga di contesto in testa."""
    if not intestazione:
        return chunks
    return [f"{intestazione}\n{c}" for c in chunks]
